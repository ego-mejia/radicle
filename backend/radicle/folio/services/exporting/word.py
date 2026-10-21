"""Word export with python-docx, written from the document's JSON: one .docx, a page per
page. Visuals come in as an image with their description when a picture of them was
taken (`pictures`), else as their description; live Radicle blocks leave a note."""

import io
from collections.abc import Callable
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from sqlalchemy.ext.asyncio import AsyncSession

from radicle.folio.services.exporting.html_render import LIVE, text_of
from radicle.folio.services.exporting.pages import ExportAsset, ExportPage, load_assets

Node = dict[str, Any]
MONO = "Consolas"
WIDTH = Inches(6)
LIVE_NOTE = "Live content from Radicle: open the page in Radicle to see it."


def _link(paragraph: Paragraph, text: str, url: str) -> None:
    """python-docx has no hyperlink API: write the run inside a w:hyperlink."""
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([style, underline])
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _inline(paragraph: Paragraph, nodes: list[Node]) -> None:
    for n in nodes:
        kind, attrs = n.get("type"), n.get("attrs") or {}
        if kind == "text":
            marks = {m.get("type"): m for m in n.get("marks") or []}
            href = str((marks.get("link") or {}).get("attrs", {}).get("href", ""))
            if href.startswith(("http://", "https://", "mailto:")):
                _link(paragraph, n.get("text", ""), href)
                continue
            run = paragraph.add_run(n.get("text", ""))
            run.bold = "bold" in marks or None
            run.italic = "italic" in marks or None
            run.underline = "underline" in marks or None
            run.font.strike = "strike" in marks or None
            if "code" in marks:
                run.font.name = MONO
        elif kind == "hardBreak":
            paragraph.add_run().add_break()
        elif kind == "mention":
            paragraph.add_run(f"@{attrs.get('handle', '')}")
        elif kind in ("status", "ticket"):
            paragraph.add_run(str(attrs.get("text") or attrs.get("key") or "")).bold = True
        elif kind == "inlineFormula":
            paragraph.add_run(str(attrs.get("latex", ""))).font.name = MONO


class _Writer:
    def __init__(
        self,
        doc: DocxDocument,
        assets: dict[str, ExportAsset],
        pictures: Callable[[Node], bytes | None],
    ) -> None:
        self.doc, self.assets, self.pictures = doc, assets, pictures

    def note(self, text: str) -> None:
        self.doc.add_paragraph().add_run(text).italic = True

    def picture(self, data: bytes, caption: str | None) -> bool:
        try:
            self.doc.add_picture(io.BytesIO(data), width=WIDTH)
        except Exception:  # a format Word can't hold (webp, svg...): keep the caption
            return False
        if caption:
            self.doc.add_paragraph(caption, style="Caption")
        return True

    def blocks(self, nodes: list[Node], depth: int = 0) -> None:
        previous: Node | None = None
        for n in nodes:
            self.block(n, depth, previous)
            previous = n

    def lists(self, items: list[Node], style: str, depth: int, mark: Callable[[Node], str]) -> None:
        name = style if depth == 0 else f"{style} {min(depth + 1, 3)}"
        for item in items:
            for i, child in enumerate(item.get("content") or []):
                if child.get("type") in ("bulletList", "orderedList", "taskList"):
                    self.block(child, depth + 1, None)
                    continue
                p = self.doc.add_paragraph(style=name if i == 0 else None)
                if i == 0 and mark(item):
                    p.add_run(mark(item))
                _inline(p, child.get("content") or [])

    def block(self, n: Node, depth: int, previous: Node | None) -> None:
        kind, attrs, content = n.get("type"), n.get("attrs") or {}, n.get("content") or []
        doc = self.doc
        if kind == "paragraph":
            _inline(doc.add_paragraph(), content)
        elif kind == "heading":
            _inline(doc.add_heading(level=min(max(int(attrs.get("level", 1)), 1), 4)), content)
        elif kind == "bulletList":
            self.lists(content, "List Bullet", depth, lambda _: "")
        elif kind == "orderedList":
            self.lists(content, "List Number", depth, lambda _: "")
        elif kind == "taskList":
            self.lists(
                content,
                "List Bullet",
                depth,
                lambda i: "☑ " if (i.get("attrs") or {}).get("checked") else "☐ ",
            )
        elif kind == "actionItem":
            p = doc.add_paragraph("☑ " if attrs.get("done") else "☐ ")
            _inline(p, content)
            if attrs.get("due"):
                p.add_run(f" ({attrs['due']})").italic = True
        elif kind in ("blockquote", "panel"):
            for child in content:
                _inline(doc.add_paragraph(style="Quote"), child.get("content") or [])
        elif kind == "codeBlock":
            run = doc.add_paragraph().add_run(text_of(n))
            run.font.name, run.font.size = MONO, Pt(9)
        elif kind == "horizontalRule":
            doc.add_paragraph("—" * 12)
        elif kind == "table":
            self.table([[text_of(c) for c in r.get("content") or []] for r in content])
        elif kind == "expand":
            doc.add_paragraph().add_run(str(attrs.get("title") or "")).bold = True
            self.blocks(content, depth)
        elif kind in ("columns", "comparison"):
            for side in content:
                title = (side.get("attrs") or {}).get("title")
                if title:
                    doc.add_paragraph().add_run(str(title)).bold = True
                self.blocks(side.get("content") or [], depth)
        elif kind == "image":
            found = self.assets.get(str(attrs.get("assetId")))
            caption = attrs.get("caption") or attrs.get("alt")
            if not (found and self.picture(found.data, caption)) and caption:
                self.note(str(caption))
        elif kind in ("file", "video"):
            found = self.assets.get(str(attrs.get("assetId")))
            name = attrs.get("name") or attrs.get("caption") or (found and found.asset.filename)
            self.note(f"📎 {name or attrs.get('src') or ''}")
        elif kind in ("visual", "diagram", "chart"):
            shot = self.pictures(n)
            description = attrs.get("description") or attrs.get("caption") or attrs.get("title")
            if shot and self.picture(shot, description):
                return
            if kind == "chart":
                rows = attrs.get("data")
                if attrs.get("source") == "table" and previous and previous.get("type") == "table":
                    rows = [
                        [text_of(c) for c in r.get("content") or []]
                        for r in previous.get("content") or []
                    ]
                if rows:
                    self.table([[str(c) for c in r] for r in rows])
            elif kind == "diagram":
                run = doc.add_paragraph().add_run(str(attrs.get("source", "")))
                run.font.name, run.font.size = MONO, Pt(9)
            if description:
                doc.add_paragraph(str(description), style="Caption")
        elif kind == "formula":
            doc.add_paragraph().add_run(str(attrs.get("latex", ""))).font.name = MONO
        elif kind == "timeline":
            for item in attrs.get("items") or []:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{item.get('date', '')} ").bold = True
                p.add_run(str(item.get("title", "")))
                if item.get("text"):
                    p.add_run(f" — {item['text']}")
        elif kind == "embed":
            url = str(attrs.get("url", ""))
            _link(doc.add_paragraph(), str(attrs.get("title") or url), url)
        elif kind in LIVE:
            self.note(LIVE_NOTE)
        elif content:
            self.blocks(content, depth)

    def table(self, rows: list[list[str]]) -> None:
        if not rows:
            return
        width = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=width)
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
                if r == 0:
                    for run in table.cell(r, c).paragraphs[0].runs:
                        run.bold = True


async def write(
    db: AsyncSession,
    pages: list[ExportPage],
    pictures: Callable[[Node], bytes | None] = lambda _: None,
) -> bytes:
    assets = await load_assets(db, pages)
    doc = Document()
    writer = _Writer(doc, assets, pictures)
    for i, page in enumerate(pages):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(page.node.title, level=0)
        description = page.version.metadata_.get("description")
        if description:
            doc.add_paragraph(str(description), style="Subtitle")
        writer.blocks(page.version.content_json.get("content") or [])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
