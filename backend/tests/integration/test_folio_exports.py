"""Exports (Folio 5.2): who may export, what an export holds, and the Markdown zip."""

import io
import uuid
import zipfile

import pytest
from docx import Document
from httpx import AsyncClient
from PIL import Image
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from radicle.folio.db.exports import ExportFormat, ExportStatus
from radicle.folio.db.models import FolioNode, FolioSpace
from radicle.folio.services import exports
from radicle.folio.services.exporting import browser
from radicle.platform.db.models import User
from radicle.platform.errors import Forbidden, NotFound, ValidationFailed
from tests.conftest import Team, run_jobs
from tests.integration.folio_helpers import API, publish


async def _user(db: AsyncSession, handle: str) -> User:
    user = await db.scalar(select(User).where(User.handle == handle))
    assert user is not None
    return user


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (4, 4), "green").save(out, "PNG")
    return out.getvalue()


@pytest.fixture
async def guides(client: AsyncClient, db: AsyncSession, team: Team) -> dict:
    """Guías: a folder with a page holding an image, a restricted page and a draft."""
    i = team.instance.headers
    await client.post(f"{API}/spaces", headers=i, json={"key": "guias", "name": "Guías"})

    async def node(title: str, kind: str = "page", parent: str | None = None) -> str:
        body = {"space": "guias", "kind": kind, "title": title, "parent": parent}
        r = await client.post(f"{API}/nodes", headers=i, json=body)
        assert r.status_code == 201, r.text
        return str(r.json()["id"])

    folder = await node("Empezar", "folder")
    guide = await node("Conectar un agente", parent=folder)
    image = await client.post(
        f"{API}/nodes/{guide}/assets", headers=i, files={"file": ("diagrama.png", _png())}
    )
    asset = image.json()["id"]
    await publish(client, team, guide, f"# Paso 1\n\n![Diagrama](asset:{asset})\n")
    secret = await node("Salarios")
    await publish(client, team, secret, "Confidencial")
    view = {"view": [{"grantee_type": "user", "grantee": team.instance.handle}]}
    assert (
        await client.put(f"{API}/nodes/{secret}/restrictions", headers=i, json=view)
    ).status_code == 200
    draft = await node("Borrador")
    grants = {
        "grants": [
            {"grantee_type": "everyone", "permissions": ["view"]},
            {
                "grantee_type": "user",
                "grantee": team.member.handle,
                "permissions": ["view", "export"],
            },
            {"grantee_type": "user", "grantee": team.instance.handle, "permissions": ["admin"]},
        ]
    }
    assert (
        await client.put(f"{API}/spaces/guias/grants", headers=i, json=grants)
    ).status_code == 200
    space = await db.scalar(select(FolioSpace).where(FolioSpace.key == "guias"))
    return {"space": space, "guide": guide, "secret": secret, "draft": draft, "asset": asset}


async def _node(db: AsyncSession, node_id: str) -> FolioNode:
    return await db.get_one(FolioNode, uuid.UUID(node_id))


async def test_a_space_to_markdown_holds_what_its_author_sees(
    db: AsyncSession, team: Team, guides: dict
) -> None:
    luis = await _user(db, team.member.handle)
    row = await exports.create_export(db, luis, guides["space"], ExportFormat.markdown)
    assert row.status == ExportStatus.queued
    await run_jobs(db)
    await db.refresh(row)
    assert row.status == ExportStatus.done, row.stats
    assert row.stats["pages"] == 1 and row.filename == "guias-markdown.zip"

    _, data = await exports.download(db, luis, row.id)
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = zf.namelist()
    page = "empezar/conectar-un-agente.md"
    assert page in names and not any("salarios" in n or "borrador" in n for n in names)
    text = zf.read(page).decode()
    assert text.startswith('---\ntitle: "Conectar un agente"\nversion: 1\n')
    assert 'doc_type: "how-to"' in text
    image = next(n for n in names if n.startswith("assets/") and n.endswith("diagrama.png"))
    assert f"](../{image}" in text and zf.read(image).startswith(b"\x89PNG")


async def test_one_page_alone(db: AsyncSession, team: Team, guides: dict) -> None:
    admin = await _user(db, team.instance.handle)
    row = await exports.create_export(
        db, admin, await _node(db, guides["guide"]), ExportFormat.markdown
    )
    await run_jobs(db)
    await db.refresh(row)
    zf = zipfile.ZipFile(io.BytesIO((await exports.download(db, admin, row.id))[1]))
    page = "conectar-un-agente.md"
    assert row.filename == "conectar-un-agente-markdown.zip" and page in zf.namelist()
    assert "](assets/" in zf.read(page).decode()


async def test_a_space_to_a_static_html_site(db: AsyncSession, team: Team, guides: dict) -> None:
    luis = await _user(db, team.member.handle)
    row = await exports.create_export(db, luis, guides["space"], ExportFormat.html)
    await run_jobs(db)
    await db.refresh(row)
    assert row.status == ExportStatus.done, row.stats
    zf = zipfile.ZipFile(io.BytesIO((await exports.download(db, luis, row.id))[1]))
    index = zf.read("index.html").decode()
    assert '<a href="empezar/conectar-un-agente.html">Conectar un agente</a>' in index
    assert "Salarios" not in index and "style.css" in zf.namelist()
    page = zf.read("empezar/conectar-un-agente.html").decode()
    assert 'href="../style.css"' in page and 'href="../index.html">← Guías' in page
    image = next(n for n in zf.namelist() if n.endswith("diagrama.png"))
    assert f'src="../{image}"' in page


async def test_a_space_to_word(db: AsyncSession, team: Team, guides: dict) -> None:
    luis = await _user(db, team.member.handle)
    row = await exports.create_export(db, luis, guides["space"], ExportFormat.docx)
    await run_jobs(db)
    await db.refresh(row)
    assert row.status == ExportStatus.done and row.filename == "guias.docx", row.stats
    doc = Document(io.BytesIO((await exports.download(db, luis, row.id))[1]))
    texts = [p.text for p in doc.paragraphs]
    assert "Conectar un agente" in texts and "Paso 1" in texts
    assert "Salarios" not in texts and len(doc.inline_shapes) == 1  # the guide's image


async def test_who_may_export_and_what(db: AsyncSession, team: Team, guides: dict) -> None:
    vera = await _user(db, team.viewer.handle)  # can read Guías, can't export
    with pytest.raises(Forbidden, match="export permission"):
        await exports.create_export(db, vera, guides["space"], ExportFormat.markdown)
    admin = await _user(db, team.instance.handle)
    with pytest.raises(ValidationFailed, match="Publish"):
        await exports.create_export(db, admin, await _node(db, guides["draft"]), ExportFormat.pdf)
    row = await exports.create_export(db, admin, guides["space"], ExportFormat.markdown)
    with pytest.raises(ValidationFailed, match="queued"):
        await exports.download(db, admin, row.id)
    luis = await _user(db, team.member.handle)
    with pytest.raises(NotFound):
        await exports.get_export(db, luis, row.id)
    assert [r.id for r in await exports.list_exports(db, admin)] == [row.id]


async def test_a_failure_is_reported_on_the_export(
    db: AsyncSession, team: Team, guides: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def no_browser(url: str) -> bytes:
        raise RuntimeError("Chromium is not installed")

    monkeypatch.setattr(browser, "pdf", no_browser)
    admin = await _user(db, team.instance.handle)
    row = await exports.create_export(db, admin, guides["space"], ExportFormat.pdf)
    await run_jobs(db)
    await db.refresh(row)
    assert row.status == ExportStatus.failed and "Chromium" in row.stats["error"]


async def test_a_pdf_prints_the_print_view_its_author_sees(
    client: AsyncClient, db: AsyncSession, team: Team, guides: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    opened: list[str] = []

    async def fake_pdf(url: str) -> bytes:
        opened.append(url)
        return b"%PDF-1.7 printed"

    monkeypatch.setattr(browser, "pdf", fake_pdf)
    luis = await _user(db, team.member.handle)
    row = await exports.create_export(db, luis, guides["space"], ExportFormat.pdf)
    await run_jobs(db)
    await db.refresh(row)
    assert row.status == ExportStatus.done and row.filename == "guias.pdf", row.stats
    assert (await exports.download(db, luis, row.id))[1].startswith(b"%PDF")
    token = opened[0].rsplit("/print/", 1)[1]

    view = await client.get(f"{API}/print/{token}")  # the browser has no account
    assert view.status_code == 200, view.text
    assert [p["title"] for p in view.json()["pages"]] == ["Conectar un agente"]
    image = await client.get(f"{API}/print/{token}/assets/{guides['asset']}")
    assert image.status_code == 200 and image.content.startswith(b"\x89PNG")
    other = await client.get(f"{API}/print/{token}/assets/00000000-0000-0000-0000-000000000000")
    assert other.status_code == 404
    assert (await client.get(f"{API}/print/{token[:-2]}xx")).status_code == 404


async def test_word_carries_a_picture_of_each_visual(
    client: AsyncClient, db: AsyncSession, team: Team, guides: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def fake_shots(codes: list[str]) -> list[bytes | None]:
        return [_png() for _ in codes]

    monkeypatch.setattr(browser, "shots", fake_shots)
    svg = '<svg viewBox="0 0 10 10"><rect width="10" height="10"/></svg>'
    md = f'```visual{{format=svg title="Flujo" description="El flujo del token"}}\n{svg}\n```\n'
    await publish(client, team, guides["guide"], md)
    admin = await _user(db, team.instance.handle)
    row = await exports.create_export(
        db, admin, await _node(db, guides["guide"]), ExportFormat.docx
    )
    await run_jobs(db)
    await db.refresh(row)
    doc = Document(io.BytesIO((await exports.download(db, admin, row.id))[1]))
    assert len(doc.inline_shapes) == 1
    assert "El flujo del token" in [p.text for p in doc.paragraphs]
