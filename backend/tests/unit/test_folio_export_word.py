"""Folio documents written into Word: every block, links, and visuals as pictures."""

import io

from docx import Document
from PIL import Image

from radicle.folio.services.exporting.word import _Writer
from tests.unit.folio_docs import EVERY_BLOCK


def _png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(out, "PNG")
    return out.getvalue()


def _text(doc: Document) -> str:
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_every_block_is_written() -> None:
    doc = Document()
    _Writer(doc, {}, lambda _: None).blocks(EVERY_BLOCK["content"])
    out = io.BytesIO()
    doc.save(out)  # a valid file
    assert len(doc.paragraphs) > 10 and doc.tables


def test_visuals_come_as_a_picture_with_their_description() -> None:
    visual = {
        "type": "visual",
        "attrs": {"format": "svg", "code": "<svg/>", "description": "El flujo del token"},
    }
    with_shot = Document()
    _Writer(with_shot, {}, lambda _: _png()).blocks([visual])
    assert len(with_shot.inline_shapes) == 1 and "El flujo del token" in _text(with_shot)
    without = Document()
    _Writer(without, {}, lambda _: None).blocks([visual])
    assert not without.inline_shapes and "El flujo del token" in _text(without)


def test_links_and_live_blocks() -> None:
    doc = Document()
    paragraph = {
        "type": "paragraph",
        "content": [
            {
                "type": "text",
                "text": "la guía",
                "marks": [{"type": "link", "attrs": {"href": "https://radicle.dev"}}],
            }
        ],
    }
    _Writer(doc, {}, lambda _: None).blocks(
        [paragraph, {"type": "pageList", "attrs": {"label": "mcp"}}]
    )
    xml = doc.paragraphs[0]._p.xml
    assert "w:hyperlink" in xml and "la guía" in xml
    assert "Live content from Radicle" in _text(doc)
    rels = [r.target_ref for r in doc.part.rels.values() if r.is_external]
    assert rels == ["https://radicle.dev"]
